import base64
import io
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skill-lab'))
from document_parser import parse_document
from document_bridge import extract
from server import Store, LabError


class DocumentTests(unittest.TestCase):
    def test_json_strict_and_original_retained(self):
        raw = '{"comments":[{"text":"英文下仍显示中文"}]}'.encode()
        result = extract('voc.json', 'application/json', raw)
        self.assertIn('英文下仍显示中文', result['text'])
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'test.sqlite3')
            case = store.create_case({'title': 'JSON', 'kind': 'voc', 'product_line': 'deebot', 'text': '',
                'attachments': [{'name': 'voc.json', 'mime': 'application/json', 'data': base64.b64encode(raw).decode()}]})
            self.assertEqual(case['attachments'][0]['parsed']['format'], 'json')
            self.assertIn('parsed', store.state()['cases'][0]['attachments'][0])
            with store.db() as db:
                self.assertEqual(base64.b64decode(store.get(db, 'cases', case['id'])['attachments'][0]['data']), raw)

    def test_bad_json_and_protected_input_rejected(self):
        for raw in (b'{"a":1,"a":2}', b'{"x":NaN}', b'not json', b'3'):
            with self.assertRaises(RuntimeError):
                parse_document('x.json', '', raw)
        with self.assertRaisesRegex(RuntimeError, '受保护'):
            parse_document('x.xlsx', '', b'%TSD-Header-test')

    def test_excel_sheets_formulas_and_zero(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = 'VOC'
        book.active.append(['ID', '反馈', '数值'])
        book.active.append(['001', '中文问题', 0])
        book.active['D2'] = '=1+2'
        book.create_sheet('指标')['A1'] = '指标数据'
        stream = io.BytesIO()
        book.save(stream)
        result = parse_document('test.xlsx', '', stream.getvalue())
        self.assertEqual(result['units'], 2)
        self.assertIn('C2=0', result['text'])
        self.assertIn('未计算', result['text'])
        self.assertIn('指标数据', result['text'])

    def test_pdf_pages_and_scanned_rejection(self):
        from reportlab.pdfgen.canvas import Canvas
        stream = io.BytesIO()
        canvas = Canvas(stream)
        canvas.drawString(20, 700, 'Evidence page one')
        canvas.showPage()
        canvas.drawString(20, 700, 'Evidence page two')
        canvas.save()
        result = parse_document('test.pdf', '', stream.getvalue())
        self.assertEqual(result['units'], 2)
        self.assertIn('[PDF 第 2 页]', result['text'])
        from pypdf import PdfWriter
        blank = PdfWriter()
        blank.add_blank_page(200, 200)
        stream = io.BytesIO()
        blank.write(stream)
        with self.assertRaisesRegex(RuntimeError, '无可提取文本'):
            parse_document('scan.pdf', '', stream.getvalue())

    def test_powerpoint_text_and_notes(self):
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = '需求目标'
        slide.notes_slide.notes_text_frame.text = '重要备注'
        stream = io.BytesIO()
        deck.save(stream)
        result = parse_document('test.pptx', '', stream.getvalue())
        self.assertIn('需求目标', result['text'])
        self.assertIn('重要备注', result['text'])
        self.assertTrue(result['warnings'])

    def test_failed_import_atomic(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'test.sqlite3')
            with self.assertRaises(LabError):
                store.create_case({'title': 'bad', 'kind': 'voc', 'product_line': 'deebot', 'text': '',
                    'attachments': [{'name': 'bad.json', 'mime': 'application/json', 'data': base64.b64encode(b'{bad').decode()}]})
            self.assertEqual(store.state()['cases'], [])
